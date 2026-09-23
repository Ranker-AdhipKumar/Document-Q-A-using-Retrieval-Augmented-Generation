"""
Knowledge Graph extraction and serialization.

Pipeline:
  1. Extract raw text from uploaded file (PDF / DOCX / TXT / MD).
  2. Call Gemini with structured output (JSON mode) to extract entities + relations.
  3. Merge into a NetworkX DiGraph (deduplicating by lowercase name).
  4. Serialize to {nodes, links} JSON for D3.js.
"""
from __future__ import annotations

import os
import re
import logging
from pathlib import Path
from typing import Optional

import networkx as nx
from google import genai
from google.genai import types

from models import (
    Entity,
    Relation,
    KnowledgeGraphExtraction,
    GraphNode,
    GraphEdge,
    GraphResponse,
)

logger = logging.getLogger(__name__)

GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.8-flash")

# Module-level graph — accumulated across all uploaded documents
_graph: nx.DiGraph = nx.DiGraph()


# ─── Text Extraction ─────────────────────────────────────────────────────────

def extract_text(file_path: str, filename: str) -> str:
    """Extract plain text from PDF, DOCX, TXT, or MD files."""
    ext = Path(filename).suffix.lower()

    if ext == ".pdf":
        return _extract_pdf(file_path)
    elif ext == ".docx":
        return _extract_docx(file_path)
    elif ext in (".txt", ".md", ".rst", ".text"):
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            return f.read()
    else:
        # Fallback: try reading as text
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                return f.read()
        except Exception:
            return ""


def _extract_pdf(file_path: str) -> str:
    try:
        import PyPDF2
        text_parts: list[str] = []
        with open(file_path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)
        return "\n".join(text_parts)
    except Exception as e:
        logger.warning("PDF extraction failed: %s", e)
        return ""


def _extract_docx(file_path: str) -> str:
    try:
        from docx import Document
        doc = Document(file_path)
        paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
        # Also extract table cell text
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
        return "\n".join(paragraphs)
    except Exception as e:
        logger.warning("DOCX extraction failed: %s", e)
        return ""


# ─── Entity Extraction via Gemini Structured Output ──────────────────────────

def _truncate_for_extraction(text: str, max_chars: int = 12_000) -> str:
    """Keep first+last portions so we capture intro and conclusion entities."""
    if len(text) <= max_chars:
        return text
    half = max_chars // 2
    return text[:half] + "\n\n[... middle truncated for extraction ...]\n\n" + text[-half:]


def extract_entities_from_text(
    text: str, document_name: str
) -> KnowledgeGraphExtraction:
    """Use Gemini structured output to extract entities and relations."""
    client = _get_client()

    truncated = _truncate_for_extraction(text)

    prompt = (
        "You are a knowledge graph builder. Analyze the following document text and extract:\n"
        "1. Important ENTITIES (people, organizations, concepts, terms, events, locations, products).\n"
        "   - Limit to the 25 most significant entities.\n"
        "   - Use canonical names (e.g., 'Machine Learning' not 'ML').\n"
        "2. RELATIONS between pairs of these entities.\n"
        "   - Limit to the 35 most meaningful relationships.\n"
        "   - Use short, lowercase labels (e.g., 'causes', 'part_of', 'created_by', 'related_to').\n\n"
        "Document text:\n"
        "---\n"
        f"{truncated}\n"
        "---\n\n"
        "Return ONLY the JSON — no markdown fences, no explanation."
    )

    import time
    models_to_try = [GEMINI_MODEL, "gemini-3.5-flash-lite"]
    for model_name in models_to_try:
        try:
            response = client.models.generate_content(
                model=model_name,
                contents=prompt,
                config={
                    "response_mime_type": "application/json",
                    "response_schema": KnowledgeGraphExtraction,
                },
            )
            if response.parsed and isinstance(response.parsed, KnowledgeGraphExtraction):
                return response.parsed
            import json
            raw = response.text or "{}"
            data = json.loads(raw)
            raw_entities = data.get("entities", []) if isinstance(data, dict) else []
            raw_relations = data.get("relations", []) if isinstance(data, dict) else []
            return KnowledgeGraphExtraction(
                entities=[Entity.from_dict(e) for e in raw_entities if isinstance(e, dict)],
                relations=[Relation.from_dict(r) for r in raw_relations if isinstance(r, dict)],
            )
        except Exception as e:
            logger.warning("Entity extraction attempt on %s failed: %s", model_name, e)
            time.sleep(1)

    return KnowledgeGraphExtraction(entities=[], relations=[])


# ─── Graph Accumulation ───────────────────────────────────────────────────────

def _normalize(name: str) -> str:
    return re.sub(r"\s+", " ", name.strip().lower())


def merge_into_graph(extraction: KnowledgeGraphExtraction, document_name: str) -> None:
    """Merge extracted entities and relations into the module-level graph."""
    global _graph

    # Map normalized name → canonical name (first-seen wins)
    for entity in extraction.entities:
        nid = _normalize(entity.name)
        if nid not in _graph:
            _graph.add_node(
                nid,
                label=entity.name,
                type=entity.type,
                description=entity.description,
                document_sources=[document_name],
            )
        else:
            # Accumulate document sources
            sources: list[str] = _graph.nodes[nid].get("document_sources", [])
            if document_name not in sources:
                sources.append(document_name)
            _graph.nodes[nid]["document_sources"] = sources

    for relation in extraction.relations:
        src = _normalize(relation.from_entity)
        tgt = _normalize(relation.to_entity)
        # Only add edge if both nodes exist (avoids phantom nodes)
        if src in _graph and tgt in _graph:
            if _graph.has_edge(src, tgt):
                _graph[src][tgt]["weight"] = _graph[src][tgt].get("weight", 1.0) + 1.0
            else:
                _graph.add_edge(src, tgt, label=relation.label, weight=1.0)


def get_graph_response() -> GraphResponse:
    """Serialize the accumulated graph to {nodes, links} for D3."""
    nodes = [
        GraphNode(
            id=nid,
            label=data.get("label", nid),
            type=data.get("type", "other"),
            description=data.get("description", ""),
            document_sources=data.get("document_sources", []),
        )
        for nid, data in _graph.nodes(data=True)
    ]

    links = [
        GraphEdge(
            source=src,
            target=tgt,
            label=data.get("label", "related_to"),
            weight=float(data.get("weight", 1.0)),
        )
        for src, tgt, data in _graph.edges(data=True)
    ]

    return GraphResponse(nodes=nodes, links=links)


def reset_graph() -> None:
    global _graph
    _graph = nx.DiGraph()


def graph_stats() -> dict:
    return {"nodes": _graph.number_of_nodes(), "edges": _graph.number_of_edges()}


# ─── Client helper ────────────────────────────────────────────────────────────

_client: genai.Client | None = None


def _get_client() -> genai.Client:
    global _client
    if _client is None:
        _client = genai.Client()
    return _client
